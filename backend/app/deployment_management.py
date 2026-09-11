from __future__ import annotations

import json
import os
import re
import threading
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader

from .azure_devops import azdo_request, find_pipeline_by_name

_DATA_FILE = Path(os.getenv("DEPLOYMENT_MANAGEMENT_FILE", "/app/data/deployment-management.json"))
_DOCUMENT_DIR = Path(os.getenv("DEPLOYMENT_DOCUMENT_DIR", "/app/data/deployment-documents"))
_LOCK = threading.RLock()

APPLICATION_TYPES = ("H2H", "Collections", "GTB-Applications", "Native-Mobile", "Safenet")
COLLECTIONS_COUNTRIES = ("Egypt", "UAE")

ARCHITECTURES: dict[str, dict[str, Any]] = {
    "H2H": {"code_pull": False, "pull_request": False, "build": True, "deployment": True},
    "Collections": {"code_pull": False, "pull_request": False, "build": True, "deployment": True, "deployment_type": "Image based"},
    "GTB-Applications": {"code_pull": True, "pull_request": True, "build": True, "deployment": True},
    "Native-Mobile": {"code_pull": False, "pull_request": False, "build": True, "deployment": True, "deployment_type": "Helm"},
    "Safenet": {"code_pull": False, "pull_request": False, "build": True, "deployment": True},
}

COLLECTIONS_PIPELINES = {
    "Egypt": {
        "vtbatch-01": "collections-batch-01-service-egypt",
        "vtbatch-02": "collections-batch-02-service-egypt",
        "vtbatch-03": "collections-batch-03-service-egypt",
        "vtbatch-04": "collections-batch-04-service-egypt",
        "vtbatch-05": "collections-batch-05-service-egypt",
        "vtchequecash": "collections-cheque-cash-service-egypt",
        "vtcustomer": "collections-customer-service-egypt",
        "vtdashboardreports": "collections-dashboard-reports-service-egypt",
        "vtdds": "collections-dds-service-egypt",
        "vtmaster": "collections-master-service-egypt",
        "vtuserrole": "collections-user-role-management-service-egypt",
        "vtransact-ui": "collections-static-web-egypt",
        "vtwidget": "collections-widget-service-egypt",
        "vtopenapi": "collections-open-api-service-egypt",
    },
    "UAE": {
        "vtbatch-01": "collections-batch-01-service",
        "vtbatch-02": "collections-batch-02-service",
        "vtbatch-03": "collections-batch-03-service",
        "vtbatch-04": "collections-batch-04-service",
        "vtbatch-05": "collections-batch-05-service",
        "vtchequecash": "collections-cheque-cash-service",
        "vtcustomer": "collections-customer-service",
        "vtdashboardreports": "collections-dashboard-reports-service",
        "vtdds": "collections-dds-service",
        "vtmaster": "collections-master-service",
        "vtuserrole": "collections-user-role-management-service",
        "vtransact-ui": "collections-static-web",
        "vtwidget": "collections-widget-service",
        "vtopenapi": "collections-open-api-service",
    },
}
COLLECTIONS_ALIASES = {"vtdashrep": "vtdashboardreports"}

REPO_MAPPING = {
    "obtf": "obtf-r2", "obtf-kernel": "obtf-r2-kernel", "obp": "obp-r2",
    "obp-kernel": "obp-r2-kernel", "obtfpm": "obtfpm-r2", "moc": "moc-r2",
    "cmncore": "commoncore-r2", "plato": "plato-r2", "obdx": "obdx-r2",
    "oblm": "oblm-r2", "oblmic": "oblm-ic-r2", "obvam": "obvam-r2", "obvamic": "obvam-ic-r2",
}
BUILD_PIPELINES = {
    "moc": "new-moc-build-pipeline", "obtfpm": "new-obtfpm-build-pipeline",
    "cmncore": "new-cmncore-build-pipeline", "plato": "new-plato-build-pipeline",
    "obtf": "r2-obtf-build-pipeline", "obp": "r2-obp-build-pipeline",
    "obdx": "r2-obdx-build-pipeline", "oblm": "new-oblm-build-pipeline",
    "oblmic": "new-oblm-ic-build-pipeline", "obvam": "new-obvam-build-pipeline",
    "obvamic": "new-obvam-ic-build-pipeline",
}
ENVIRONMENT_BRANCH = {
    "R2": "release/r2", "R2UAT": "release/uat", "UAT": "release/uat",
    "PREPRD": "release/ppr", "PPR": "release/ppr", "R2TRAIN": "release/train",
    "PROD": "release/prod", "GOLD": "release/gold", "SIT": "release/sit",
}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _parse_time(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None


def _duration_seconds(started_at: str | None, finished_at: str | None) -> int | None:
    started = _parse_time(started_at)
    finished = _parse_time(finished_at)
    if not started:
        return None
    end = finished or datetime.now(timezone.utc)
    return max(0, int((end - started).total_seconds()))


def _load() -> list[dict[str, Any]]:
    with _LOCK:
        if not _DATA_FILE.exists():
            return []
        try:
            value = json.loads(_DATA_FILE.read_text(encoding="utf-8"))
            return value if isinstance(value, list) else []
        except (OSError, json.JSONDecodeError):
            return []


def _save(rows: list[dict[str, Any]]) -> None:
    with _LOCK:
        _DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
        temp = _DATA_FILE.with_suffix(".tmp")
        temp.write_text(json.dumps(rows, indent=2), encoding="utf-8")
        temp.replace(_DATA_FILE)


def _next_deployment_request_id(rows: list[dict[str, Any]]) -> str:
    date_key = datetime.now(timezone.utc).strftime("%Y%m%d")
    prefix = f"DM-{date_key}-"
    highest = 0
    for row in rows:
        request_id = str(row.get("id") or "")
        if not request_id.startswith(prefix):
            continue
        suffix = request_id[len(prefix):]
        if suffix.isdigit():
            highest = max(highest, int(suffix))
    return f"{prefix}{highest + 1:04d}"


def _extract_text(name: str, raw: bytes) -> str:
    lowered = name.lower()
    if lowered.endswith(".pdf"):
        reader = PdfReader(BytesIO(raw))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if lowered.endswith(".docx"):
        document = Document(BytesIO(raw))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if lowered.endswith((".txt", ".md")):
        return raw.decode("utf-8", errors="replace")
    raise HTTPException(status_code=400, detail="Supported release documents are PDF, DOCX, TXT and MD")


def _first_match(patterns: list[str], text: str) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
        if match:
            return match.group(1).strip().strip(".,;:")
    return ""


def _collections_items(text: str, country: str) -> list[dict[str, Any]]:
    mapping = COLLECTIONS_PIPELINES.get(country, {})
    seen: set[str] = set()
    items: list[dict[str, Any]] = []
    pattern = r"[A-Za-z0-9.-]+\.azurecr\.io/([A-Za-z0-9._/-]+):([A-Za-z0-9_.-]+)"
    for image_name, tag in re.findall(pattern, text):
        extracted_service = image_name.rsplit("/", 1)[-1].strip().lower()
        service = COLLECTIONS_ALIASES.get(extracted_service, extracted_service)
        if service in seen:
            continue
        seen.add(service)
        items.append({
            "selected": True,
            "service": service,
            "image_tag": tag,
            "vendor_image": f"{extracted_service}:{tag}",
            "use_vendor_image": True,
            "extraction_method": "Rule-Based",
            "pipeline_name": mapping.get(service, ""),
            "country": country,
        })
    return items


def _structured_extract(filename: str, raw: bytes, application_type: str = "", country: str = "") -> dict[str, Any]:
    text = _extract_text(filename, raw)
    branch = _first_match([
        r"(?:source\s+|profinch\s+)?branch(?:\s+name)?\s*[:=-]\s*([^\s,;|]+)",
        r"\b((?:release|feature|hotfix|bugfix)/[A-Za-z0-9._/-]+)\b",
    ], text)
    application = _first_match([
        r"(?:pipeline\s+)?application(?:\s+name)?\s*[:=-]\s*([A-Za-z0-9._-]+)",
        r"service(?:\s+name)?\s*[:=-]\s*([A-Za-z0-9._-]+)",
    ], text)
    environment = _first_match([
        r"environment\s*[:=-]\s*([A-Za-z0-9_-]+)",
        r"\b(R2UAT|PREPRD|PPR|PROD|R2TRAIN|GOLD|SIT|UAT|R2)\b",
    ], text).upper()
    repository = _first_match([r"repository(?:\s+name)?\s*[:=-]\s*([A-Za-z0-9._-]+)"], text)
    war_files = sorted(set(re.findall(r"\b[\w.-]+\.war\b", text, flags=re.IGNORECASE)))
    jar_files = sorted(set(re.findall(r"\b[\w.-]+\.jar\b", text, flags=re.IGNORECASE)))
    images = sorted(set(re.findall(r"[\w.-]+\.azurecr\.io/[\w./-]+:[\w.-]+", text, flags=re.IGNORECASE)))
    app_key = application.strip().lower()
    result = {
        "filename": filename,
        "application_type": application_type,
        "application": application,
        "branch_name": branch,
        "environment": environment,
        "repository": repository or REPO_MAPPING.get(app_key, application),
        "target_branch": ENVIRONMENT_BRANCH.get(environment, "release/uat"),
        "build_pipeline": BUILD_PIPELINES.get(app_key, ""),
        "war_files": war_files,
        "jar_files": jar_files,
        "container_images": images,
        "text_preview": text[:8000],
        "confidence": {name: ("high" if value else "missing") for name, value in {"application": application, "branch_name": branch, "environment": environment}.items()},
    }
    if application_type == "Collections":
        result["country"] = country
        result["collections_items"] = _collections_items(text, country)
        result["application"] = "Collections"
        result["environment"] = country
    return result


def extract_release_document(upload: UploadFile, raw: bytes, application_type: str = "", country: str = "") -> dict[str, Any]:
    if application_type and application_type not in APPLICATION_TYPES:
        raise HTTPException(status_code=422, detail="Unsupported application type")
    return _structured_extract(upload.filename or "release-document", raw, application_type, country)


def submit_document(application_type: str, app_owner: str, upload: UploadFile, raw: bytes, requested_by: str, country: str = "") -> dict[str, Any]:
    if application_type not in APPLICATION_TYPES:
        raise HTTPException(status_code=422, detail="Select a valid application type")
    if application_type == "Collections" and country not in COLLECTIONS_COUNTRIES:
        raise HTTPException(status_code=422, detail="Collections country must be Egypt or UAE")
    if not app_owner.strip():
        raise HTTPException(status_code=422, detail="Application owner is required")
    if not raw:
        raise HTTPException(status_code=422, detail="Release document is empty")
    now = _now()
    rows = _load()
    request_id = _next_deployment_request_id(rows)
    _DOCUMENT_DIR.mkdir(parents=True, exist_ok=True)
    suffix = Path(upload.filename or "release-document").suffix.lower()
    document_path = _DOCUMENT_DIR / f"{request_id}{suffix}"
    document_path.write_bytes(raw)
    architecture = ARCHITECTURES[application_type]
    row = {
        "id": request_id, "status": "Pending App Owner Approval", "requested_by": requested_by,
        "created_at": now, "updated_at": now, "application_type": application_type,
        "country": country, "app_owner": app_owner.strip(), "document_name": upload.filename or "release-document",
        "document_path": str(document_path), "extracted": False, "architecture": architecture,
        "application": "", "branch_name": "", "environment": country if application_type == "Collections" else "",
        "repository": "", "target_branch": "", "build_pipeline": "", "war_files": [], "jar_files": [],
        "container_images": [], "collections_items": [], "progress_percent": 10,
        "steps": {
            "owner_approval": {"status": "Pending"}, "devops_review": {"status": "Waiting"},
            "document_extraction": {"status": "Waiting"},
            "code_pull": {"status": "Not Required" if not architecture.get("code_pull") else "Waiting"},
            "pull_request": {"status": "Not Required" if not architecture.get("pull_request") else "Waiting"},
            "build": {"status": "Waiting", "runs": []}, "deployment": {"status": "Waiting", "runs": []},
        },
        "timeline": [{"at": now, "action": f"Document submitted{f' for {country}' if country else ''}", "actor": requested_by}],
    }
    rows.insert(0, row); _save(rows)
    return row


def create_deployment_request(payload: dict[str, Any], requested_by: str) -> dict[str, Any]:
    raise HTTPException(status_code=410, detail="Developers must submit a release document using the document submission endpoint")


def list_deployment_requests() -> list[dict[str, Any]]:
    return _load()


def _pipeline_run(pipeline_name: str, parameters: dict[str, Any], pat: str, ref: str = "") -> dict[str, Any]:
    pipeline = find_pipeline_by_name(pipeline_name, pat)
    if not pipeline:
        raise HTTPException(status_code=404, detail=f"Pipeline {pipeline_name} was not found")
    body: dict[str, Any] = {"templateParameters": parameters}
    normalized_ref = ref.strip()
    if normalized_ref:
        if not normalized_ref.startswith("refs/heads/"):
            normalized_ref = f"refs/heads/{normalized_ref}"
        body["resources"] = {"repositories": {"self": {"refName": normalized_ref}}}
    code, run = azdo_request("POST", f"_apis/pipelines/{pipeline['id']}/runs?api-version=7.1", pat, body)
    if code not in (200, 201):
        raise HTTPException(status_code=502, detail=run.get("message", f"Unable to trigger {pipeline_name}"))
    run_id = run.get("id")
    web_url = ((run.get("_links") or {}).get("web") or {}).get("href")
    return {
        "status": "Queued",
        "result": "",
        "pipeline_name": pipeline_name,
        "pipeline_id": pipeline["id"],
        "run_id": run_id,
        "url": web_url,
        "logs_url": web_url,
        "started_at": run.get("createdDate") or _now(),
        "finished_at": None,
        "duration_seconds": None,
    }


def _refresh_pipeline_run(run: dict[str, Any], pat: str) -> dict[str, Any]:
    pipeline_id = run.get("pipeline_id")
    run_id = run.get("run_id")
    if not pipeline_id or not run_id:
        return run
    code, body = azdo_request("GET", f"_apis/pipelines/{pipeline_id}/runs/{run_id}?api-version=7.1", pat)
    if code != 200:
        run["error"] = body.get("message", f"Unable to refresh pipeline run {run_id}")
        return run
    state = str(body.get("state") or "").strip()
    result = str(body.get("result") or "").strip()
    started_at = body.get("createdDate") or body.get("startedDate") or run.get("started_at")
    finished_at = body.get("finishedDate") or run.get("finished_at")
    web_url = ((body.get("_links") or {}).get("web") or {}).get("href") or run.get("url")
    normalized_status = "Running"
    if state.lower() in {"notstarted", "queued"}:
        normalized_status = "Queued"
    elif state.lower() in {"completed"}:
        normalized_status = "Succeeded" if result.lower() == "succeeded" else "Failed"
    run.update({
        "status": normalized_status,
        "result": result,
        "url": web_url,
        "logs_url": web_url,
        "started_at": started_at,
        "finished_at": finished_at,
        "duration_seconds": _duration_seconds(started_at, finished_at),
    })
    return run


def _find_release_for_build(pipeline_name: str, build_run_id: int, pat: str) -> dict[str, Any] | None:
    code, definitions = azdo_request("GET", "_apis/release/definitions?$top=500&api-version=7.1", pat)
    if code != 200:
        return None
    candidates = definitions.get("value", [])
    exact = next((d for d in candidates if str(d.get("name") or "").casefold() == pipeline_name.casefold()), None)
    definition = exact or next((d for d in candidates if pipeline_name.casefold() in str(d.get("name") or "").casefold()), None)
    if not definition:
        return None
    code, releases = azdo_request("GET", f"_apis/release/releases?definitionId={definition.get('id')}&$top=20&queryOrder=descending&api-version=7.1", pat)
    if code != 200:
        return None
    fallback = None
    for release in releases.get("value", []):
        release_id = release.get("id")
        code, detail = azdo_request("GET", f"_apis/release/releases/{release_id}?api-version=7.1", pat)
        if code != 200:
            continue
        if fallback is None:
            fallback = detail
        for artifact in detail.get("artifacts", []):
            instance = artifact.get("instanceReference") or {}
            definition_ref = artifact.get("definitionReference") or {}
            version_ref = definition_ref.get("version") or {}
            if str(instance.get("id")) == str(build_run_id) or str(version_ref.get("id")) == str(build_run_id):
                return detail
    return fallback


def _refresh_collections_release_links(row: dict[str, Any], pat: str) -> None:
    build_runs = (row.get("steps") or {}).get("build", {}).get("runs", [])
    deployment_runs: list[dict[str, Any]] = []
    for build_run in build_runs:
        if str(build_run.get("status")) != "Succeeded":
            continue
        release = _find_release_for_build(str(build_run.get("pipeline_name") or ""), int(build_run.get("run_id") or 0), pat)
        if not release:
            continue
        web_url = ((release.get("_links") or {}).get("web") or {}).get("href")
        environments = release.get("environments", [])
        status = "Running"
        environment_states = [str(env.get("status") or "") for env in environments]
        if environment_states and all(state.lower() in {"succeeded", "rejected", "canceled", "cancelled"} for state in environment_states):
            status = "Succeeded" if all(state.lower() == "succeeded" for state in environment_states) else "Failed"
        deployment_runs.append({
            "pipeline_name": release.get("name") or build_run.get("pipeline_name"),
            "release_id": release.get("id"),
            "status": status,
            "result": status,
            "url": web_url,
            "logs_url": web_url,
            "started_at": release.get("createdOn"),
            "finished_at": release.get("modifiedOn") if status in {"Succeeded", "Failed"} else None,
            "duration_seconds": _duration_seconds(release.get("createdOn"), release.get("modifiedOn") if status in {"Succeeded", "Failed"} else None),
            "environment_status": ", ".join(f"{env.get('name')}={env.get('status')}" for env in environments),
        })
    if deployment_runs:
        row.setdefault("steps", {}).setdefault("deployment", {})["runs"] = deployment_runs
        statuses = [run.get("status") for run in deployment_runs]
        if all(status == "Succeeded" for status in statuses):
            row["steps"]["deployment"]["status"] = "Succeeded"
            row["status"] = "Deployment Completed"
        elif any(status == "Failed" for status in statuses):
            row["steps"]["deployment"]["status"] = "Failed"
            row["status"] = "Deployment Failed"
        else:
            row["steps"]["deployment"]["status"] = "Running"
            row["status"] = "Deployment Running"


def _update_progress(row: dict[str, Any]) -> None:
    if row.get("application_type") != "Collections":
        return
    progress = 10
    if row.get("steps", {}).get("owner_approval", {}).get("status") == "Approved":
        progress = 25
    if row.get("extracted"):
        progress = 40
    build_status = row.get("steps", {}).get("build", {}).get("status")
    if build_status in {"Queued", "Running"}:
        progress = 60
    elif build_status == "Succeeded":
        progress = 80
    elif build_status == "Failed":
        progress = 60
    deployment_status = row.get("steps", {}).get("deployment", {}).get("status")
    if deployment_status in {"Queued", "Running"}:
        progress = 90
    elif deployment_status == "Succeeded":
        progress = 100
    elif deployment_status == "Failed":
        progress = 90
    row["progress_percent"] = progress


def refresh_request_status(request_id: str, pat: str) -> dict[str, Any]:
    rows = _load()
    row = next((item for item in rows if item.get("id") == request_id), None)
    if not row:
        raise HTTPException(status_code=404, detail="Deployment request not found")
    steps = row.setdefault("steps", {})
    build_runs = steps.setdefault("build", {}).setdefault("runs", [])
    for run in build_runs:
        _refresh_pipeline_run(run, pat)
    if build_runs:
        statuses = [str(run.get("status") or "") for run in build_runs]
        if all(status == "Succeeded" for status in statuses):
            steps["build"]["status"] = "Succeeded"
            row["status"] = "Build Completed"
        elif any(status == "Failed" for status in statuses):
            steps["build"]["status"] = "Failed"
            row["status"] = "Build Failed"
        elif any(status == "Running" for status in statuses):
            steps["build"]["status"] = "Running"
            row["status"] = "Build Running"
        else:
            steps["build"]["status"] = "Queued"
            row["status"] = "Build Queued"
    deployment_runs = steps.setdefault("deployment", {}).setdefault("runs", [])
    for run in deployment_runs:
        if run.get("pipeline_id") and run.get("run_id"):
            _refresh_pipeline_run(run, pat)
    if row.get("application_type") == "Collections" and steps.get("build", {}).get("status") == "Succeeded":
        _refresh_collections_release_links(row, pat)
    _update_progress(row)
    now = _now()
    row["updated_at"] = now
    row.setdefault("timeline", []).append({"at": now, "action": "refresh-status", "actor": "System"})
    _save(rows)
    return row


def _trigger_collections_builds(row: dict[str, Any], updates: dict[str, Any], pat: str) -> dict[str, Any]:
    selected = updates.get("collections_items") or row.get("collections_items") or []
    selected = [item for item in selected if item.get("selected", True)]
    if not selected:
        raise HTTPException(status_code=422, detail="Select at least one Collections image")
    runs = []
    for item in selected:
        pipeline_name = item.get("pipeline_name") or COLLECTIONS_PIPELINES.get(row.get("country", ""), {}).get(item.get("service", ""), "")
        if not pipeline_name:
            runs.append({**item, "status": "Failed", "result": "mapping_missing", "error": "Pipeline mapping not found"})
            continue
        use_vendor_image = bool(item.get("use_vendor_image", True))
        vendor_image = str(item.get("vendor_image") or "").strip()
        if use_vendor_image and not vendor_image:
            runs.append({**item, "status": "Failed", "result": "invalid_parameters", "error": "vendorImage is required when useVendorImage is enabled"})
            continue
        parameters = {
            "vendorImage": vendor_image,
            "useVendorImage": use_vendor_image,
        }
        run = _pipeline_run(pipeline_name, parameters, pat)
        runs.append({**item, **run, "use_vendor_image": use_vendor_image, "parameters": parameters})
    row["collections_items"] = selected
    overall = "Queued" if runs and all(run.get("status") == "Queued" for run in runs) else "Running"
    return {"status": overall, "runs": runs}


def update_action(request_id: str, action: str, actor: str, pat: str = "", updates: dict[str, Any] | None = None) -> dict[str, Any]:
    rows = _load(); row = next((item for item in rows if item.get("id") == request_id), None)
    if not row:
        raise HTTPException(status_code=404, detail="Deployment request not found")
    now = _now(); steps = row.setdefault("steps", {}); updates = updates or {}
    if action == "owner-approve":
        row["status"] = "Awaiting DevOps"; steps["owner_approval"] = {"status": "Approved", "actor": actor, "at": now}
    elif action == "owner-reject":
        row["status"] = "App Owner Rejected"; steps["owner_approval"] = {"status": "Rejected", "actor": actor, "at": now}
    elif action == "extract-document":
        path = Path(row["document_path"]); extracted = _structured_extract(row["document_name"], path.read_bytes(), row["application_type"], row.get("country", ""))
        row.update(extracted); row["extracted"] = True; row["status"] = "DevOps Review"
        steps["document_extraction"] = {"status": "Completed", "actor": actor, "at": now}
        steps["devops_review"] = {"status": "In Progress", "actor": actor, "at": now}
    elif action == "save-extracted-data":
        for key in ("application", "branch_name", "environment", "repository", "target_branch", "build_pipeline", "war_files", "jar_files", "container_images", "collections_items"):
            if key in updates: row[key] = updates[key]
        row["status"] = "Ready for Orchestration"; steps["devops_review"] = {"status": "Completed", "actor": actor, "at": now}
    elif action == "trigger-code-pull":
        run = _pipeline_run(os.getenv("CODE_PULL_PIPELINE_NAME", "code-pull-pipeline"), {
            os.getenv("CODE_PULL_PARAM_APPLICATION", "APP"): row["application"],
            os.getenv("CODE_PULL_PARAM_BRANCH", "PROFINCH_BRANCH"): row["branch_name"],
            os.getenv("CODE_PULL_PARAM_LIST_ONLY", "LIST_ONLY"): False,
        }, pat, os.getenv("CODE_PULL_PIPELINE_REF", "refs/heads/master"))
        row["status"] = "Code Pull Running"; steps["code_pull"] = run
    elif action == "create-pr":
        source = updates.get("source_branch") or row.get("generated_branch") or row["branch_name"]
        payload = {"sourceRefName": f"refs/heads/{source}", "targetRefName": f"refs/heads/{row['target_branch']}",
                   "title": f"{row['application']} deployment {row['environment']}", "description": f"Deployment request {row['id']}"}
        code, pr = azdo_request("POST", f"_apis/git/repositories/{row['repository']}/pullrequests?api-version=7.1", pat, payload)
        if code not in (200, 201): raise HTTPException(status_code=502, detail=pr.get("message", "Unable to create PR"))
        steps["pull_request"] = {"status": "Active", "pr_id": pr.get("pullRequestId"), "url": ((pr.get("_links") or {}).get("web") or {}).get("href")}
        row["status"] = "PR Awaiting Review"
    elif action == "trigger-build":
        if row.get("application_type") == "Collections":
            steps["build"] = _trigger_collections_builds(row, updates, pat)
            row["status"] = "Build Queued"
        else:
            pipeline_name = row.get("build_pipeline") or updates.get("pipeline_name")
            if not pipeline_name: raise HTTPException(status_code=422, detail="Build pipeline is not configured")
            parameters = updates.get("parameters") or {"branch": row.get("target_branch"), "war_files": " ".join(row.get("war_files") or []) or "None", "jar_files": " ".join(row.get("jar_files") or []) or "None", "Environment": row.get("environment")}
            steps["build"] = {"status": "Queued", "runs": [_pipeline_run(pipeline_name, parameters, pat)]}; row["status"] = "Build Queued"
    elif action == "trigger-deployment":
        if row.get("application_type") == "Collections":
            if steps.get("build", {}).get("status") != "Succeeded":
                raise HTTPException(status_code=409, detail="All Collections build pipelines must succeed before deployment")
            _refresh_collections_release_links(row, pat)
            if not steps.get("deployment", {}).get("runs"):
                pipeline_name = updates.get("pipeline_name") or os.getenv(f"COLLECTIONS_DEPLOY_PIPELINE_{row.get('country', '').upper()}", "")
                if not pipeline_name:
                    raise HTTPException(status_code=422, detail="No release was discovered and Collections deployment pipeline name is not configured")
                run = _pipeline_run(pipeline_name, updates.get("parameters") or {"country": row.get("country"), "environment": row.get("environment"), "services": ",".join(item.get("service", "") for item in row.get("collections_items", []))}, pat)
                steps["deployment"] = {"status": "Queued", "runs": [run]}
                row["status"] = "Deployment Queued"
        else:
            pipeline_name = updates.get("pipeline_name") or os.getenv(f"{row['application_type'].upper().replace('-', '_')}_DEPLOY_PIPELINE", "")
            if not pipeline_name: raise HTTPException(status_code=422, detail="Deployment pipeline name is required")
            run = _pipeline_run(pipeline_name, updates.get("parameters") or {"environment": row.get("environment"), "application": row.get("application")}, pat)
            steps["deployment"] = {"status": "Queued", "runs": [run]}; row["status"] = "Deployment Queued"
    elif action == "refresh-status":
        return refresh_request_status(request_id, pat)
    else:
        raise HTTPException(status_code=400, detail="Unsupported deployment action")
    _update_progress(row)
    row["updated_at"] = now
    row.setdefault("timeline", []).append({"at": now, "action": action, "actor": actor})
    _save(rows)
    return row
